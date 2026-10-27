import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook

from academics.models import College, Department, Year
from data_import.services import ImportRejected, run_import
from faculty.models import Faculty


def _make_workbook(faculty_rows=None, subject_rows=None, classroom_rows=None, omit_sheet=None,
                    faculty_headers=None):
    wb = Workbook()
    wb.remove(wb.active)

    if omit_sheet != "Faculty":
        ws = wb.create_sheet("Faculty")
        headers = faculty_headers or [
            "faculty_name", "email", "username", "designation", "short_code",
            "max_lecture_per_day", "max_lecture_per_week",
        ]
        ws.append(headers)
        for row in (faculty_rows or [["Dr. A", "a@test.edu", "", "Asst Prof", "DRA", 6, 24]]):
            ws.append(row)

    if omit_sheet != "Subject":
        ws = wb.create_sheet("Subject")
        ws.append(["sr.no", "course type", "subject code", "subject name", "lecture", "tutorial", "practical", "credits"])
        for row in (subject_rows or [[1, "CORE", "CS301", "DBMS", 3, 0, 2, 4.5]]):
            ws.append(row)

    if omit_sheet != "Classroom":
        ws = wb.create_sheet("Classroom")
        ws.append(["sr.no", "branch name and section", "classroom number", "building name"])
        for row in (classroom_rows or [[1, "CSE - A", "CR-501", "Main Block"]]):
            ws.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _upload(buf, name="import.xlsx"):
    return SimpleUploadedFile(name, buf.read(), content_type="application/vnd.openxmlformats")


class ExcelImportTests(TestCase):
    def setUp(self):
        college = College.objects.create(name="T", code="T1")
        self.dept = Department.objects.create(college=college, code="CSE", name="CS")
        self.year = Year.objects.create(department=self.dept, year_number=2, is_active=True)

    def test_valid_import_succeeds(self):
        f = _upload(_make_workbook(faculty_rows=[["Dr. A", "a@test.edu", "dra", "Asst Prof", "DRA", 6, 24]]))
        batch = run_import(f, department=self.dept, year=self.year, uploaded_by=None)
        self.assertEqual(batch.status, "SUCCESS")
        self.assertEqual(batch.faculty_imported, 1)
        self.assertEqual(batch.subjects_imported, 1)
        self.assertEqual(batch.classrooms_imported, 1)
        self.assertTrue(Faculty.objects.filter(email="a@test.edu").exists())

    def test_username_derived_when_blank_and_warned(self):
        f = _upload(_make_workbook(faculty_rows=[["Dr. B", "bob.smith@test.edu", "", "", "BSM", 6, 24]]))
        batch = run_import(f, department=self.dept, year=self.year, uploaded_by=None)
        self.assertEqual(batch.status, "PARTIAL")
        self.assertTrue(any("derived" in w for w in batch.warnings))

    def test_missing_required_sheet_rejects_whole_import(self):
        f = _upload(_make_workbook(omit_sheet="Classroom"))
        batch = run_import(f, department=self.dept, year=self.year, uploaded_by=None)
        self.assertEqual(batch.status, "FAILED")
        self.assertEqual(batch.faculty_imported, 0)
        self.assertFalse(Faculty.objects.exists())
        self.assertTrue(any("Classroom" in e for e in batch.errors))

    def test_missing_required_header_rejects_whole_import(self):
        f = _upload(_make_workbook(faculty_headers=["faculty_name", "email"]))  # missing several required headers
        batch = run_import(f, department=self.dept, year=self.year, uploaded_by=None)
        self.assertEqual(batch.status, "FAILED")
        self.assertFalse(Faculty.objects.exists())

    def test_duplicate_faculty_within_file_rejected(self):
        rows = [
            ["Dr. A", "dup@test.edu", "", "", "DUP", 6, 24],
            ["Dr. A2", "dup@test.edu", "", "", "DU2", 6, 24],
        ]
        f = _upload(_make_workbook(faculty_rows=rows))
        batch = run_import(f, department=self.dept, year=self.year, uploaded_by=None)
        self.assertEqual(batch.status, "FAILED")
        self.assertFalse(Faculty.objects.exists())  # whole import rolled back

    def test_invalid_file_type_rejected(self):
        f = SimpleUploadedFile("not_excel.txt", b"hello", content_type="text/plain")
        batch = run_import(f, department=self.dept, year=self.year, uploaded_by=None)
        self.assertEqual(batch.status, "FAILED")
        self.assertTrue(any(".txt" in e or "file type" in e for e in batch.errors))

    def test_odd_practical_hours_rejected(self):
        f = _upload(_make_workbook(subject_rows=[[1, "CORE", "CS999", "Weird Lab", 2, 0, 3, 4]]))
        batch = run_import(f, department=self.dept, year=self.year, uploaded_by=None)
        self.assertEqual(batch.status, "FAILED")
        self.assertTrue(any("even" in e for e in batch.errors))

    def test_reimport_updates_existing_records_idempotently(self):
        f1 = _upload(_make_workbook())
        run_import(f1, department=self.dept, year=self.year, uploaded_by=None)
        self.assertEqual(Faculty.objects.count(), 1)

        f2 = _upload(_make_workbook(faculty_rows=[["Dr. A Updated", "a@test.edu", "", "Professor", "DRA", 5, 20]]))
        run_import(f2, department=self.dept, year=self.year, uploaded_by=None)
        self.assertEqual(Faculty.objects.count(), 1)
        fac = Faculty.objects.get(email="a@test.edu")
        self.assertEqual(fac.name, "Dr. A Updated")
        self.assertEqual(fac.max_lectures_per_day, 5)
