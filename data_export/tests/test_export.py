from django.test import TestCase
from openpyxl import load_workbook

from common.constants import GenerationStatus
from data_export.services import export_excel, export_ical, export_pdf
from scheduling.models import TimetableVersion
from scheduling.services.generator import run_generation
from scheduling.tests.factories import build_minimal_scope
from scheduling.tests.test_generation import _create_generation


class ExportTests(TestCase):
    def setUp(self):
        self.scope = build_minimal_scope()
        generation = _create_generation(self.scope)
        run_generation(generation)
        self.assertEqual(generation.status, GenerationStatus.CURRENT)
        self.version = TimetableVersion.objects.get(section=self.scope["section"], status=GenerationStatus.CURRENT)

    def test_excel_export_contains_subject_codes(self):
        buf = export_excel(self.version, college_name="Test College")
        wb = load_workbook(buf)
        ws = wb["Timetable"]
        all_text = "\n".join(
            str(cell.value) for row in ws.iter_rows() for cell in row if cell.value
        )
        self.assertIn("CS301", all_text)
        self.assertIn("CS302", all_text)
        self.assertIn("Test College", all_text)

    def test_pdf_export_produces_nonempty_pdf(self):
        buf = export_pdf(self.version, college_name="Test College")
        data = buf.read()
        self.assertTrue(data.startswith(b"%PDF"))
        self.assertGreater(len(data), 1000)

    def test_ical_export_produces_valid_calendar_with_events(self):
        buf = export_ical(self.version)
        from icalendar import Calendar
        cal = Calendar.from_ical(buf.read())
        events = [c for c in cal.walk() if c.name == "VEVENT"]
        self.assertGreater(len(events), 0)
        summaries = {str(e.get("summary")) for e in events}
        self.assertTrue(any("CS301" in s for s in summaries))
        self.assertTrue(any("CS302" in s for s in summaries))
